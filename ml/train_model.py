"""
Train the E-Optimizer consumption model and export it for the browser.

Inputs : process levers from Grain_Based_Ethanol_Synthetic_Dataset_2025.xlsx
Outputs: Total_Process_Electricity_kWh, Distillation_Steam_kg, DDGS_Dryer_Fuel_MMBtu

The trained network is written as plain JSON (weights, biases, scaler stats) so
the frontend can run the forward pass directly with no ML runtime. An sklearn
MLPRegressor forward pass is just matmul + ReLU, so this costs nothing in
accuracy and saves shipping a wasm runtime.

Deliberately NOT predicted:
  Ethanol_Production_kL      - equals Grain_Input_tpd * 0.39 exactly, so it is a
                               formula, not something to learn.
  Ethanol_Yield_L_per_t_Grain- constant 390.0 in this dataset, zero variance.

Run:  python ml/train_model.py
"""

import json
import sys
from pathlib import Path

import numpy as np
import openpyxl
from sklearn.model_selection import train_test_split
from sklearn.neural_network import MLPRegressor
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parent.parent
DATASET = ROOT.parent.parent / "Grain_Based_Ethanol_Synthetic_Dataset_2025.xlsx"
OUT_PATH = ROOT / "WEB" / "public" / "model" / "ann.json"

# Every candidate lever considered during feature selection.
CANDIDATE_FEATURES = [
    "Grain_Input_tpd",
    "Grain_Moisture_pct",
    "Grain_Starch_pct",
    "Milling_Particle_Size_micron",
    "Milling_Efficiency_pct",
    "Liquefaction_Temp_C",
    "Liquefaction_pH",
    "Liquefaction_Time_min",
    "Slurry_Solids_pct",
    "Alpha_Amylase_Dosage_kg_per_t",
    "Saccharification_Temp_C",
    "Saccharification_pH",
    "Saccharification_Time_hr",
    "Glucoamylase_Dosage_kg_per_t",
    "Fermentation_Temp_C",
    "Fermentation_pH",
    "Fermentation_Time_hr",
    "Yeast_Dosage_kg_per_t",
    "Distillation_Temp_C",
    "Reflux_Ratio",
    "Dryer_Inlet_Temp_C",
    "Condensate_Recovery_pct",
    "DDGS_Moisture_pct",
]

# Greedy forward selection under 5-fold CV chose throughput alone. No second
# lever improved held-out accuracy, and using all 23 made every target worse
# (electricity 0.60 -> 0.49, steam 0.27 -> 0.19, fuel 0.13 -> 0.00). Adding the
# rest fits noise: outside grain input, no lever exceeds |r| = 0.19 against any
# target in this dataset.
FEATURES = ["Grain_Input_tpd"]

TARGETS = [
    "Total_Process_Electricity_kWh",
    "Distillation_Steam_kg",
    "DDGS_Dryer_Fuel_MMBtu",
]

RANDOM_STATE = 42


def load_dataset(path: Path):
    if not path.exists():
        sys.exit(f"Dataset not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Daily_Data_2025"]
    header = [c.value for c in ws[1]]
    index = {name: i for i, name in enumerate(header)}

    missing = [c for c in FEATURES + TARGETS if c not in index]
    if missing:
        sys.exit(f"Dataset is missing expected columns: {missing}")

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    X = np.array([[r[index[c]] for c in FEATURES] for r in rows], dtype=float)
    y = np.array([[r[index[c]] for c in TARGETS] for r in rows], dtype=float)
    return X, y


def r2(actual: np.ndarray, predicted: np.ndarray) -> float:
    ss_res = float(((actual - predicted) ** 2).sum())
    ss_tot = float(((actual - actual.mean()) ** 2).sum())
    return 1.0 - ss_res / ss_tot


def main() -> None:
    X, y = load_dataset(DATASET)
    print(f"Loaded {X.shape[0]} rows, {X.shape[1]} features, {y.shape[1]} targets")

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=RANDOM_STATE
    )

    x_scaler = StandardScaler().fit(X_train)
    y_scaler = StandardScaler().fit(y_train)

    # Small and strongly regularised on purpose: 365 rows with one informative
    # feature. A bigger net just memorises noise.
    model = MLPRegressor(
        hidden_layer_sizes=(16,),
        activation="relu",
        solver="adam",
        alpha=1.0,
        max_iter=8000,
        early_stopping=True,
        n_iter_no_change=80,
        random_state=RANDOM_STATE,
    )
    model.fit(x_scaler.transform(X_train), y_scaler.transform(y_train))

    predicted = y_scaler.inverse_transform(model.predict(x_scaler.transform(X_test)))

    print("\nHeld-out test R2 (20% split):")
    scores = {}
    for i, name in enumerate(TARGETS):
        score = r2(y_test[:, i], predicted[:, i])
        scores[name] = round(score, 4)
        print(f"  {name:32s} {score:7.4f}")

    print(
        "\nCeiling is the dataset, not the model. Grain_Input_tpd correlates "
        "0.81/0.58/0.43 with the three targets; every other lever sits below "
        "|r| = 0.19, which is noise at n=365."
    )

    payload = {
        "features": FEATURES,
        "targets": TARGETS,
        "xMean": x_scaler.mean_.tolist(),
        "xScale": x_scaler.scale_.tolist(),
        "yMean": y_scaler.mean_.tolist(),
        "yScale": y_scaler.scale_.tolist(),
        "weights": [w.tolist() for w in model.coefs_],
        "biases": [b.tolist() for b in model.intercepts_],
        "activation": "relu",
        "testR2": scores,
        "trainedRows": int(X.shape[0]),
        "candidatesConsidered": len(CANDIDATE_FEATURES),
        "selectionMethod": "greedy forward selection, 5-fold CV",
        "note": (
            "Trained on a synthetic dataset. Its README states it is not for "
            "regulatory reporting or carbon-credit verification."
        ),
    }

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(payload), encoding="utf-8")
    size_kb = OUT_PATH.stat().st_size / 1024
    print(f"\nWrote {OUT_PATH.relative_to(ROOT)} ({size_kb:.1f} kB)")


if __name__ == "__main__":
    main()
